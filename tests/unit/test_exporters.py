import os
import tempfile
import pytest
from pptx import Presentation

from backend.app.exporters.markdown_exporter import MarkdownExporter
from backend.app.exporters.pptx_exporter import PPTXExporter
from backend.app.exporters.pdf_exporter import PDFExporter
from backend.app.exporters.manager import ExportManager


@pytest.fixture
def sample_state():
    return {
        "idea": "AI tutor for Class 10 students in India",
        "status": "completed",
        "scores": {
            "market": [8.0],
            "prd": [7.5],
            "architecture": [8.5],
            "pitch": [9.0]
        },
        "market": {
            "tam": {"value": "$50B", "source_url": "https://example.com"},
            "sam": {"value": "$10B", "source_url": "https://example.com"},
            "som": {"value": "$500M", "source_url": "https://example.com"},
            "competitors": [
                {"name": "Byju's", "description": "Edtech incumbent", "strengths": ["Scale"], "weaknesses": ["Expensive"]}
            ],
            "trends": ["AI personalized learning"],
            "risks": ["Curriculum changes"],
            "sources": ["https://example.com"]
        },
        "prd": {
            "problem_statement": "Students lack affordable personalized exam coaching.",
            "mvp_scope": "6 weeks for 2 developers",
            "personas": [
                {"name": "Aarav", "role": "Student", "pain_points": ["Math struggles"], "goals": ["Score > 90%"]}
            ],
            "features": [
                {"name": "Instant Doubts AI", "priority": "P0", "description": "Solves math questions step by step."}
            ],
            "user_stories": ["As a student, I want instant explanations."],
            "success_metrics": ["70% D30 retention"]
        },
        "architecture": {
            "stack": [
                {"category": "Backend", "technology": "FastAPI", "justification": "Fast async API."}
            ],
            "diagram_mermaid": "flowchart TD\n  Client --> API",
            "monthly_cost_estimate": {
                "total_usd": 65.0,
                "breakdown": {"Compute": 30.0, "Database": 25.0, "Redis": 10.0}
            },
            "scalability_notes": "Stateless horizontal scaling.",
            "risks": ["LLM API rate limits"]
        },
        "pitch": {
            "slides": [
                {
                    "slide_number": i,
                    "title": f"Slide {i}: Strategy",
                    "bullets": ["Traction point A", "Market opportunity B"],
                    "speaker_notes": "Here is the key founder message."
                }
                for i in range(1, 11)
            ]
        },
        "debate_log": [
            {
                "round": 1,
                "from_agent": "market_analyst",
                "to_agent": "product_manager",
                "type": "critique",
                "content": {"issues": ["Add more STEM focus"]}
            }
        ]
    }


def test_markdown_export(sample_state):
    exporter = MarkdownExporter()
    content = exporter.generate("run-123", sample_state)
    assert "# Startup Package: AI tutor" in content
    assert "## 2. Market Analysis" in content
    assert "## 3. Product Requirements Document (PRD)" in content
    assert "## 4. Technical Architecture" in content
    assert "## 5. Pitch Deck" in content
    assert "## 6. Peer Debate & Critique Transcript" in content


def test_pptx_export(sample_state):
    exporter = PPTXExporter()
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "test_pitch.pptx")
        exporter.generate("run-123", sample_state, out_path)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0

        # Verify PPTX structure
        prs = Presentation(out_path)
        assert len(prs.slides) == 10
        # Check notes
        assert prs.slides[0].notes_slide.notes_text_frame.text != ""


def test_pdf_export(sample_state):
    exporter = PDFExporter()
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "test_report.pdf")
        exporter.generate("run-123", sample_state, out_path)
        assert os.path.exists(out_path)
        assert os.path.getsize(out_path) > 0


def test_export_manager(sample_state):
    with tempfile.TemporaryDirectory() as tmpdir:
        manager = ExportManager(exports_dir=tmpdir)
        results = manager.export("run-999", sample_state, file_format="all")
        assert "md" in results and os.path.exists(results["md"])
        assert "pptx" in results and os.path.exists(results["pptx"])
        assert "pdf" in results and os.path.exists(results["pdf"])
        assert "csv" in results and os.path.exists(results["csv"])


def test_csv_export(sample_state):
    from backend.app.exporters.csv_exporter import CSVExporter
    exporter = CSVExporter()
    with tempfile.TemporaryDirectory() as tmpdir:
        out_path = os.path.join(tmpdir, "test_data.csv")
        exporter.generate("run-123", sample_state, out_path)
        assert os.path.exists(out_path)
        with open(out_path, "r", encoding="utf-8-sig") as f:
            content = f.read()
        assert "=== VENTURE SUMMARY & METADATA ===" in content
        assert "=== 1. QUALITY GATE EVALUATIONS ===" in content
        assert "=== 2. MARKET SIZING (TAM / SAM / SOM) ===" in content
        assert "Total Addressable Market (TAM),$50B" in content
        assert "=== 4. PRODUCT FEATURE BACKLOG (PRD) ===" in content
        assert "Instant Doubts AI,Solves math questions step by step." in content
        assert "=== 7. CLOUD INFRASTRUCTURE RUN RATE ===" in content
        assert "Compute,$30.00" in content


def test_export_manager_in_memory(sample_state):
    import io
    manager = ExportManager()
    for fmt in ["csv", "md", "pptx", "pdf"]:
        stream = manager.export_in_memory("run-in-mem", sample_state, file_format=fmt)
        assert isinstance(stream, io.BytesIO)
        content_bytes = stream.getvalue()
        assert len(content_bytes) > 0
        if fmt == "csv":
            text = content_bytes.decode("utf-8-sig")
            assert "=== VENTURE SUMMARY & METADATA ===" in text
        elif fmt == "md":
            text = content_bytes.decode("utf-8")
            assert "# Startup Package" in text


def test_csv_formula_injection_defense():
    """Verify that CSV formula injection triggers ('=', '+', '-', '@', '\t', '\r') are neutralized."""
    from backend.app.exporters.csv_exporter import CSVExporter
    exporter = CSVExporter()
    malicious_state = {
        "idea": "=cmd|'/C calc'!A0",
        "status": "+DDE_PAYLOAD",
        "market": {
            "tam": {"value": "-@SUM(1+1)", "source_url": "https://safe.example.com"},
            "competitors": [{"name": "@EvilCorp", "description": "\tTabTrigger", "strengths": ["=1+1"], "weaknesses": ["-2"]}]
        },
        "prd": {
            "features": [{"name": "=HYPERLINK(\"http://evil.com\")", "priority": "P0", "description": "+steal"}]
        }
    }
    stream = exporter.generate("run-sec-1", malicious_state)
    content = stream.getvalue().decode("utf-8-sig")
    # All formula triggers must be prepended with a single quote
    assert "'=cmd|'/C calc'!A0" in content
    assert "'+DDE_PAYLOAD" in content
    assert "'-@SUM(1+1)" in content
    assert "'@EvilCorp" in content
    assert "'\tTabTrigger" in content
    assert "'=HYPERLINK" in content


def test_pdf_xml_escaping_and_special_chars():
    """Verify that unclosed tags, inequality symbols, ampersands, and SSRF img tags don't crash ReportLab."""
    exporter = PDFExporter()
    malicious_state = {
        "idea": "B2B SaaS with <unclosed_tag & latency < 100ms targeting <50 employees",
        "market": {
            "tam": {"value": "<img src='http://127.0.0.1:9999/fake.png'/>", "source_url": "https://example.com"},
            "competitors": [{"name": "R&D <Tech> & Co", "description": "Formula <A & B>"}],
            "trends": ["Growth > 50%"],
            "risks": ["Risk < 0"]
        },
        "prd": {
            "problem_statement": "Problem with <angle brackets> & symbols",
            "mvp_scope": "Scope <= 8 weeks",
            "features": [{"name": "Feature <P0> & Fast", "priority": "P0"}]
        },
        "architecture": {
            "stack": [{"category": "Backend", "technology": "FastAPI <async> & uv"}],
            "monthly_cost_estimate": {"total_usd": 50.0}
        },
        "pitch": {
            "slides": [{"slide_number": 1, "title": "Slide <1> & Overview"}]
        }
    }
    # This must generate successfully without raising paraparser.ParagraphParseError / ValueError
    stream = exporter.generate("run-sec-pdf", malicious_state)
    assert stream is not None
    pdf_bytes = stream.getvalue()
    assert len(pdf_bytes) > 0
    # PDF header check
    assert pdf_bytes.startswith(b"%PDF")


def test_path_traversal_defense_in_export_manager(sample_state):
    """Verify that malicious run_id with path traversal sequences is blocked."""
    with tempfile.TemporaryDirectory() as tmpdir:
        manager = ExportManager(exports_dir=tmpdir)
        # Slashes and traversal dots must be rejected
        with pytest.raises(ValueError, match="Path traversal detected"):
            manager.export("../../etc/passwd", sample_state)

        with pytest.raises(ValueError, match="Path traversal detected"):
            manager.export("..\\..\\windows\\win.ini", sample_state)


def test_markdown_table_cell_and_url_sanitization():
    """Verify that pipe delimiters in table cells are escaped and dangerous URLs are neutralized."""
    exporter = MarkdownExporter()
    state = {
        "idea": "Startup | Pipe Injection Test",
        "market": {
            "tam": {"value": "$10B", "source_url": "javascript:alert('XSS')"},
            "competitors": [
                {"name": "Comp | Split", "description": "Line1\nLine2", "strengths": ["A | B"], "weaknesses": ["None"]}
            ],
            "sources": ["javascript:stealToken()", "https://trusted-site.com"]
        },
        "prd": {
            "features": [{"name": "Auth | SSO", "priority": "P0", "description": "Single | Sign-on"}]
        },
        "architecture": {
            "stack": [{"category": "DB | Storage", "technology": "Postgres | Redis", "justification": "Fast | Scalable"}]
        }
    }
    md_content = exporter.generate("run-sec-md", state)
    # Dangerous javascript: URLs must be neutralized to '#'
    assert "javascript:alert" not in md_content
    assert "javascript:stealToken" not in md_content
    assert "https://trusted-site.com" in md_content
    # Pipes inside cells must be escaped
    assert "Comp \\| Split" in md_content
    assert "Auth \\| SSO" in md_content
    assert "DB \\| Storage" in md_content


