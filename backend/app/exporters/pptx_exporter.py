import os
from typing import Dict, Any, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


class PPTXExporter:
    """Exports Pitch Deck deliverable to professional 16:9 PowerPoint presentation (FR-10)."""

    DARK_BG = RGBColor(15, 23, 42)       # Slate 900
    TEXT_LIGHT = RGBColor(248, 250, 252) # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184) # Slate 400
    ACCENT = RGBColor(99, 102, 241)      # Indigo 500

    def generate(self, run_id: str, state: Dict[str, Any], output_path: Any = None) -> Any:
        """Create PowerPoint file or return in-memory BytesIO stream."""
        import io
        prs = Presentation()
        # Set 16:9 widescreen dimensions
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        idea = state.get("idea", "Startup Pitch Deck")
        pitch = state.get("pitch") or {}
        slides_data = pitch.get("slides") or []

        # If no slides were produced, construct default fallback slides
        if not slides_data:
            slides_data = [
                {
                    "slide_number": 1,
                    "title": idea,
                    "bullets": ["AI Startup Co-Founder Team Deck", "Automated Venture Validation"],
                    "speaker_notes": "Welcome to our investor presentation."
                }
            ]

        # Use blank slide layout
        blank_layout = prs.slide_layouts[6]

        for s_idx, slide_info in enumerate(slides_data):
            slide = prs.slides.add_slide(blank_layout)

            # Apply dark background
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = self.DARK_BG

            title_text = slide_info.get("title", f"Slide {s_idx + 1}")
            bullets = slide_info.get("bullets", [])
            speaker_notes = slide_info.get("speaker_notes", "")

            # 1. Slide Number Badge
            badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(2.0), Inches(0.4))
            badge_tf = badge_box.text_frame
            badge_tf.word_wrap = True
            p_badge = badge_tf.paragraphs[0]
            p_badge.text = f"SLIDE {s_idx + 1:02d} / {len(slides_data):02d}"
            p_badge.font.size = Pt(10)
            p_badge.font.bold = True
            p_badge.font.color.rgb = self.ACCENT

            # 2. Slide Title
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.7), Inches(1.2))
            title_tf = title_box.text_frame
            title_tf.word_wrap = True
            p_title = title_tf.paragraphs[0]
            p_title.text = title_text
            p_title.font.size = Pt(28)
            p_title.font.bold = True
            p_title.font.color.rgb = self.TEXT_LIGHT

            # 3. Bullets Body
            body_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.7), Inches(4.2))
            body_tf = body_box.text_frame
            body_tf.word_wrap = True

            for b_idx, bullet in enumerate(bullets):
                p = body_tf.paragraphs[0] if b_idx == 0 else body_tf.add_paragraph()
                p.text = f"•  {bullet}"
                p.font.size = Pt(18)
                p.font.color.rgb = self.TEXT_LIGHT
                p.space_after = Pt(14)

            # 4. Speaker Notes
            if speaker_notes:
                notes_slide = slide.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = speaker_notes

        # Return BytesIO buffer or save to file
        if output_path is None or isinstance(output_path, io.BytesIO):
            buffer = output_path if isinstance(output_path, io.BytesIO) else io.BytesIO()
            prs.save(buffer)
            buffer.seek(0)
            return buffer

        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        prs.save(output_path)
        return output_path
